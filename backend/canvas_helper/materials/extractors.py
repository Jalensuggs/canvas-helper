from __future__ import annotations

from dataclasses import dataclass, field
import html
import json
import mimetypes
from pathlib import Path
from typing import Any


@dataclass(slots=True)
class ExtractionResult:
    text: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)
    preview_html: str | None = None
    error: str | None = None


def _plain_preview(title: str, sections: list[tuple[str, str]]) -> str:
    body = "".join(
        f"<section><h2>{html.escape(heading)}</h2><pre>{html.escape(value)}</pre></section>"
        for heading, value in sections
        if value
    )
    return (
        "<!doctype html><meta charset=utf-8>"
        "<meta http-equiv=Content-Security-Policy "
        'content="default-src \'none\'; style-src \'unsafe-inline\'">'
        "<style>body{font:15px system-ui;margin:2rem;line-height:1.55;color:#20251f}"
        "pre{white-space:pre-wrap;font:inherit}section{border-bottom:1px solid #ddd}"
        "table{border-collapse:collapse}td,th{border:1px solid #ddd;padding:.4rem}</style>"
        f"<title>{html.escape(title)}</title><h1>{html.escape(title)}</h1>{body}"
    )


def _extract_pdf(path: Path) -> ExtractionResult:
    try:
        import pymupdf as fitz  # type: ignore[import-not-found]
    except ImportError:
        return ExtractionResult(error="PDF text extraction unavailable (install PyMuPDF)")
    with fitz.open(path) as document:
        return ExtractionResult(
            # Form-feed is retained for stable page-aware search chunks.
            text="\f".join(page.get_text() for page in document),
            metadata={"pages": document.page_count, **dict(document.metadata or {})},
        )


def _extract_docx(path: Path) -> ExtractionResult:
    try:
        from docx import Document as DocxDocument  # type: ignore[import-not-found]
    except ImportError:
        return ExtractionResult(error="DOCX extraction unavailable (install python-docx)")
    document = DocxDocument(path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    text = "\n".join(paragraphs)
    props = document.core_properties
    return ExtractionResult(
        text=text,
        metadata={"author": props.author, "subject": props.subject},
        preview_html=_plain_preview(path.name, [("Document", text)]),
    )


def _extract_pptx(path: Path) -> ExtractionResult:
    try:
        from pptx import Presentation  # type: ignore[import-not-found]
    except ImportError:
        return ExtractionResult(error="PPTX extraction unavailable (install python-pptx)")
    presentation = Presentation(path)
    sections: list[tuple[str, str]] = []
    for index, slide in enumerate(presentation.slides, 1):
        values = [
            shape.text for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        sections.append((f"Slide {index}", "\n".join(values)))
    return ExtractionResult(
        # Form-feed is retained for stable slide-aware search chunks.
        text="\f".join(value for _, value in sections),
        metadata={"slides": len(presentation.slides)},
        preview_html=_plain_preview(path.name, sections),
    )


def _extract_xlsx(path: Path) -> ExtractionResult:
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]
    except ImportError:
        return ExtractionResult(error="XLSX extraction unavailable (install openpyxl)")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[tuple[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                "\t".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            sections.append((sheet.title, "\n".join(rows)))
    finally:
        workbook.close()
    return ExtractionResult(
        text="\n\n".join(value for _, value in sections),
        metadata={"sheets": [name for name, _ in sections]},
        preview_html=_plain_preview(path.name, sections),
    )


def _extract_ipynb(path: Path) -> ExtractionResult:
    notebook = json.loads(path.read_text(encoding="utf-8"))
    sections: list[tuple[str, str]] = []
    for index, cell in enumerate(notebook.get("cells", []), 1):
        kind = str(cell.get("cell_type", "cell"))
        source = cell.get("source", "")
        value = "".join(source) if isinstance(source, list) else str(source)
        sections.append((f"{kind.title()} {index}", value))
    return ExtractionResult(
        text="\n\n".join(value for _, value in sections),
        metadata={
            "cells": len(sections),
            "kernel": (notebook.get("metadata", {}).get("kernelspec") or {}).get("name"),
        },
        preview_html=_plain_preview(path.name, sections),
    )


def _extract_image(path: Path) -> ExtractionResult:
    try:
        from PIL import Image  # type: ignore[import-not-found]
    except ImportError:
        return ExtractionResult(error="Image metadata unavailable (install Pillow)")
    with Image.open(path) as image:
        return ExtractionResult(
            metadata={
                "width": image.width,
                "height": image.height,
                "format": image.format,
                "mode": image.mode,
            }
        )


_OFFICE_DOCX = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}
_OFFICE_PPTX = {
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/vnd.ms-powerpoint",
}
_OFFICE_XLSX = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}


def extract_document(path: Path, mime_type: str | None = None) -> ExtractionResult:
    """Extract passive text/metadata. Optional format dependencies fail softly."""
    mime = (mime_type or mimetypes.guess_type(path.name)[0] or "").lower()
    suffix = path.suffix.lower()
    try:
        if mime == "application/pdf" or suffix == ".pdf":
            return _extract_pdf(path)
        if suffix == ".docx" or mime in _OFFICE_DOCX:
            return _extract_docx(path)
        if suffix == ".pptx" or mime in _OFFICE_PPTX:
            return _extract_pptx(path)
        if suffix == ".xlsx" or mime in _OFFICE_XLSX:
            return _extract_xlsx(path)
        if suffix == ".ipynb" or mime == "application/x-ipynb+json":
            return _extract_ipynb(path)
        if mime.startswith("image/"):
            return _extract_image(path)
        if mime.startswith("text/") or suffix in {
            ".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml", ".py",
            ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".rs",
            ".go", ".sql", ".sh", ".css",
        }:
            return ExtractionResult(text=path.read_text(encoding="utf-8", errors="replace"))
        return ExtractionResult(error="No safe extractor is available for this file type")
    except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
        return ExtractionResult(error=f"{type(exc).__name__}: {exc}")
