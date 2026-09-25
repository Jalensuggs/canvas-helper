import hashlib
import json
import sqlite3
from pathlib import Path

import httpx
import pytest
from fastapi import Request
from fastapi.testclient import TestClient

from canvas_helper.canvas.client import CanvasClient
from canvas_helper.config import Settings
from canvas_helper.main import create_app
from canvas_helper.materials import MaterialStore, extract_document
from canvas_helper.models import LOCAL_USER_ID, User
from canvas_helper.tenancy import get_current_user


@pytest.mark.asyncio
async def test_download_redirect_hash_and_user_edit_preservation(tmp_path: Path):
    payloads = [b"first version", b"second version"]
    downloads = 0

    async def handler(request: httpx.Request) -> httpx.Response:
        nonlocal downloads
        if request.url.host == "canvas.example.edu":
            assert request.headers["Authorization"] == "Bearer secret-value"
            return httpx.Response(302, headers={"Location": "https://cdn.example.edu/file"})
        assert request.url.host == "cdn.example.edu"
        assert "Authorization" not in request.headers
        payload = payloads[min(downloads, len(payloads) - 1)]
        downloads += 1
        return httpx.Response(200, content=payload, headers={"Content-Type": "text/plain"})

    canvas = CanvasClient(
        "https://canvas.example.edu",
        "secret-value",
        transport=httpx.MockTransport(handler),
        max_retries=0,
    )
    store = MaterialStore(tmp_path / "materials", LOCAL_USER_ID)
    try:
        first = await store.download(
            canvas,
            url="https://canvas.example.edu/files/7/download",
            course_id=1,
            file_id="7",
            filename="notes.txt",
        )
        assert first.sha256 == hashlib.sha256(payloads[0]).hexdigest()
        first.path.write_text("my local notes", encoding="utf-8")
        second = await store.download(
            canvas,
            url="https://canvas.example.edu/files/7/download",
            course_id=1,
            file_id="7",
            filename="notes.txt",
            prior_sha256=first.sha256,
            prior_version=first.version,
        )
    finally:
        await canvas.close()

    assert second.version == 2
    assert second.path.read_bytes() == payloads[1]
    assert second.preserved_path is not None
    assert second.preserved_path.read_text(encoding="utf-8") == "my local notes"
    assert second.preserved_path.name.startswith("user-")
    assert not list(second.path.parent.glob("*.part"))


@pytest.mark.asyncio
async def test_download_limits_https_redirects(tmp_path: Path):
    async def handler(request: httpx.Request) -> httpx.Response:
        number = int(request.url.path.rsplit("/", 1)[-1])
        return httpx.Response(
            302, headers={"Location": f"https://canvas.example.edu/r/{number + 1}"}
        )

    canvas = CanvasClient(
        "https://canvas.example.edu",
        "secret-value",
        transport=httpx.MockTransport(handler),
        max_retries=0,
    )
    try:
        with pytest.raises(httpx.TooManyRedirects):
            await MaterialStore(tmp_path / "materials", LOCAL_USER_ID).download(
                canvas,
                url="https://canvas.example.edu/r/0",
                course_id=1,
                file_id="8",
                filename="loop.txt",
            )
    finally:
        await canvas.close()


def test_material_paths_block_traversal_and_symlink(tmp_path: Path):
    store = MaterialStore(tmp_path / "materials", LOCAL_USER_ID)
    with pytest.raises(ValueError):
        store.destination(1, "../escape", "file.txt")
    with pytest.raises(ValueError):
        store.resolve_local_path("../escape")
    course = store.root / "course-1"
    course.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    (course / "file-9").symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError):
        store.destination(1, "9", "file.txt")


def test_text_and_notebook_extraction_generate_passive_preview(tmp_path: Path):
    text_path = tmp_path / "sample.py"
    text_path.write_text("print('safe')", encoding="utf-8")
    assert extract_document(text_path).text == "print('safe')"

    notebook_path = tmp_path / "lesson.ipynb"
    notebook_path.write_text(
        json.dumps(
            {
                "metadata": {"kernelspec": {"name": "python3"}},
                "cells": [{"cell_type": "code", "source": ["x = 1\n", "x"]}],
            }
        ),
        encoding="utf-8",
    )
    result = extract_document(notebook_path)
    assert "x = 1" in result.text
    assert result.metadata["cells"] == 1
    assert result.preview_html is not None
    assert "default-src 'none'" in result.preview_html


def test_material_api_tenant_boundary_and_active_content_denial(tmp_path: Path):
    path = tmp_path / "materials-api.db"
    app = create_app(
        Settings(
            database_url=f"sqlite+aiosqlite:///{path}",
            data_dir=tmp_path / "data",
        )
    )

    async def selected_user(request: Request) -> User:
        user_id = request.headers.get("X-Test-User", LOCAL_USER_ID)
        return User(id=user_id, display_name=user_id)

    app.dependency_overrides[get_current_user] = selected_user
    with TestClient(app) as client:
        with sqlite3.connect(path) as connection:
            connection.execute(
                "INSERT INTO users (id, display_name, created_at) "
                "VALUES ('tenant-two', 'Tenant Two', CURRENT_TIMESTAMP)"
            )
            document_id = connection.execute(
                "INSERT INTO documents "
                "(user_id, title, content, mime_type, kind, version, metadata_json, updated_at) "
                "VALUES ('tenant-two', 'unsafe.html', '<b>source</b>', 'text/html', "
                "'file', 1, '{}', CURRENT_TIMESTAMP) RETURNING id"
            ).fetchone()[0]
            connection.commit()

        assert client.get(f"/api/materials/{document_id}").status_code == 404
        headers = {"X-Test-User": "tenant-two"}
        assert client.get(f"/api/materials/{document_id}", headers=headers).status_code == 200
        denied = client.get(f"/api/materials/{document_id}/preview", headers=headers)
        assert denied.status_code == 415
        content = client.get(f"/api/materials/{document_id}/content", headers=headers)
        assert content.status_code == 200
        assert content.headers["content-type"].startswith("text/plain")


def test_material_browser_hides_image_assets(tmp_path: Path):
    path = tmp_path / "materials-list.db"
    app = create_app(
        Settings(
            database_url=f"sqlite+aiosqlite:///{path}",
            data_dir=tmp_path / "data",
        )
    )
    with TestClient(app) as client:
        with sqlite3.connect(path) as connection:
            course_id = connection.execute(
                "INSERT INTO courses "
                "(user_id, canvas_id, name, raw, synced_at) "
                "VALUES (?, 101, 'Current Course', '{}', CURRENT_TIMESTAMP) RETURNING id",
                (LOCAL_USER_ID,),
            ).fetchone()[0]
            for title, mime in (
                ("slides.pdf", "application/pdf"),
                ("decorative-banner.png", None),
                ("camera upload", "image/jpeg"),
            ):
                connection.execute(
                    "INSERT INTO documents "
                    "(user_id, course_id, title, content, mime_type, kind, version, "
                    "metadata_json, updated_at) VALUES (?, ?, ?, '', ?, 'file', 1, "
                    "'{}', CURRENT_TIMESTAMP)",
                    (LOCAL_USER_ID, course_id, title, mime),
                )
            connection.commit()

        response = client.get("/api/materials")
        assert response.status_code == 200
        assert [item["title"] for item in response.json()] == ["slides.pdf"]


def test_extract_document_uses_office_mime_without_suffix(tmp_path: Path):
    from docx import Document as DocxDocument

    source = tmp_path / "notes.docx"
    document = DocxDocument()
    document.add_paragraph("Hello preview")
    document.save(source)
    nameless = tmp_path / "file-9"
    nameless.write_bytes(source.read_bytes())
    result = extract_document(
        nameless,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "Hello preview" in result.text
    assert result.preview_html
    assert result.error is None


def test_material_content_reextracts_failed_office_document(tmp_path: Path):
    from docx import Document as DocxDocument

    path = tmp_path / "materials-refresh.db"
    settings = Settings(
        database_url=f"sqlite+aiosqlite:///{path}",
        data_dir=tmp_path / "data",
        allow_development_secret_file=True,
    )
    app = create_app(settings)
    source = tmp_path / "assignment.docx"
    document = DocxDocument()
    document.add_paragraph("Reflection task brief")
    document.save(source)
    with TestClient(app) as client:
        store = MaterialStore(settings.data_dir / "materials", LOCAL_USER_ID)
        stored = store.destination(1, "77", "assignment.docx")
        stored.parent.mkdir(parents=True, exist_ok=True)
        stored.write_bytes(source.read_bytes())
        with sqlite3.connect(path) as connection:
            document_id = connection.execute(
                "INSERT INTO documents "
                "(user_id, title, content, mime_type, kind, local_path, version, "
                "metadata_json, updated_at) VALUES (?, ?, '', "
                "'application/vnd.openxmlformats-officedocument.wordprocessingml.document', "
                "'file', ?, 1, ?, CURRENT_TIMESTAMP) RETURNING id",
                (
                    LOCAL_USER_ID,
                    "assignment.docx",
                    store.relative(stored),
                    '{"extraction":{"error":"DOCX extraction unavailable (install python-docx)"}}',
                ),
            ).fetchone()[0]
            connection.commit()
        content = client.get(f"/api/materials/{document_id}/content")
        assert content.status_code == 200
        assert "Reflection task brief" in content.text
        preview = client.get(f"/api/materials/{document_id}/preview")
        assert preview.status_code == 200
        assert "Reflection task brief" in preview.text
